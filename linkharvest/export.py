"""导出层：把抓取结果写成 docx / markdown / xlsx / json，并单独输出失败清单。"""

from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Item, Options

CJK_FONT = "微软雅黑"
MONO_FONT = "Consolas"


# --------------------------------------------------------------------------- 公共


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _safe_filename(text: str, fallback: str = "document") -> str:
    text = re.sub(r"[\\/:*?\"<>|\r\n\t]+", "_", text).strip(" ._")
    return text[:60] or fallback


def failed_items(items: list[Item]) -> list[Item]:
    return [it for it in items if it.status != "ok"]


# --------------------------------------------------------------------------- docx


class DocxWriter:
    def __init__(self, image_dir: Path, opts: Options) -> None:
        self.image_dir = image_dir
        self.opts = opts

    def _set_base_font(self, doc: Document) -> None:
        style = doc.styles["Normal"]
        style.font.name = CJK_FONT
        style.font.size = Pt(10.5)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.4

    def _heading(self, doc: Document, text: str, level: int) -> None:
        para = doc.add_heading(level=level)
        run = para.add_run(text)
        run.font.name = CJK_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
        run.bold = True

    def _meta_line(self, doc: Document, label: str, value: str) -> None:
        if not value:
            return
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(1)
        run = para.add_run(f"{label}：")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        val = para.add_run(value)
        val.font.size = Pt(9)
        val.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    def build(self, items: list[Item], path: Path, title: str = "链接归档") -> Path:
        doc = Document()
        self._set_base_font(doc)
        section = doc.sections[0]
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)

        doc.add_heading(title, level=0)
        ok_count = sum(1 for it in items if it.status == "ok")
        intro = doc.add_paragraph()
        intro.add_run(
            f"生成时间：{_stamp()}　共 {len(items)} 条链接，成功 {ok_count} 条，"
            f"异常 {len(items) - ok_count} 条。"
        ).font.size = Pt(9)

        # 目录表格
        self._heading(doc, "目录", 1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        table.autofit = False
        widths = [Cm(1.2), Cm(9.6), Cm(3.0), Cm(2.4)]
        headers = ["#", "标题", "时间", "状态"]
        for cell, text, width in zip(table.rows[0].cells, headers, widths):
            cell.width = width
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.size = Pt(9)
        for idx, item in enumerate(items, 1):
            row = table.add_row().cells
            values = [
                str(idx),
                item.title or item.url,
                item.date or "—",
                "OK" if item.status == "ok" else ("部分" if item.status == "partial" else "失败"),
            ]
            for cell, text, width in zip(row, values, widths):
                cell.width = width
                cell.text = text
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(1)
                    for run in para.runs:
                        run.font.size = Pt(9)

        # 正文
        for idx, item in enumerate(items, 1):
            doc.add_page_break()
            self._heading(doc, f"{idx}. {item.title or item.url}", 1)
            self._meta_line(doc, "链接", item.url)
            self._meta_line(doc, "时间", item.date or "未识别")
            self._meta_line(doc, "作者", item.author)
            self._meta_line(doc, "站点", item.sitename)
            if item.summary:
                self._meta_line(doc, "摘要", item.summary)
            if item.error:
                para = doc.add_paragraph()
                run = para.add_run(f"⚠ {item.error}")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)

            if item.content_html:
                self._render_html(doc, item.content_html)
            elif item.status == "failed":
                doc.add_paragraph("（抓取失败，未获取到正文）")

        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        return path

    # ---- 极简 HTML → docx 渲染 ----
    def _render_html(self, doc: Document, body_html: str) -> None:
        soup = BeautifulSoup(body_html, "lxml")
        root = soup.body or soup
        for node in root.children:
            if isinstance(node, NavigableString):
                text = str(node).strip()
                if text:
                    doc.add_paragraph(text)
                continue
            if not isinstance(node, Tag):
                continue
            self._render_tag(doc, node)

    def _render_tag(self, doc: Document, tag: Tag) -> None:
        name = tag.name
        if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            text = tag.get_text(" ", strip=True)
            if text:
                level = min(3, max(2, int(name[1])))
                self._heading(doc, text, level)
        elif name == "p":
            self._render_inline_paragraph(doc, tag)
        elif name in {"ul", "ol"}:
            for li in tag.find_all("li", recursive=False):
                text = li.get_text(" ", strip=True)
                if text:
                    para = doc.add_paragraph(text, style="List Bullet" if name == "ul" else "List Number")
                    para.paragraph_format.space_after = Pt(2)
        elif name == "blockquote":
            children = tag.find_all("p", recursive=False)
            texts = [c.get_text(" ", strip=True) for c in children] if children else [
                tag.get_text(" ", strip=True)
            ]
            for text in [t for t in texts if t]:
                para = doc.add_paragraph(text, style="Intense Quote")
                para.paragraph_format.space_after = Pt(4)
        elif name == "pre":
            text = tag.get_text("\n", strip=False)
            if text.strip():
                para = doc.add_paragraph()
                run = para.add_run(text.strip())
                run.font.name = MONO_FONT
                run.font.size = Pt(8.5)
                para.paragraph_format.left_indent = Cm(0.5)
        elif name == "figure":
            self._render_children(doc, tag)
        elif name == "figcaption":
            text = tag.get_text(" ", strip=True)
            if text:
                para = doc.add_paragraph(text)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(6)
                for run in para.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        elif name == "img":
            self._add_image(doc, tag)
        elif name == "table":
            self._render_table(doc, tag)
        elif name in {"div", "section", "article", "main", "span", "a", "body", "figure", "picture"}:
            self._render_children(doc, tag)
        else:
            text = tag.get_text(" ", strip=True)
            if text:
                doc.add_paragraph(text)

    def _render_children(self, doc: Document, tag: Tag) -> None:
        for child in tag.children:
            if isinstance(child, Tag):
                self._render_tag(doc, child)
            elif isinstance(child, NavigableString) and str(child).strip():
                doc.add_paragraph(str(child).strip())

    def _render_inline_paragraph(self, doc: Document, tag: Tag) -> None:
        # 段落里可能夹着图片：先出文字，再出图
        text = tag.get_text(" ", strip=True)
        if text:
            para = doc.add_paragraph(text)
            para.paragraph_format.space_after = Pt(6)
        for img in tag.find_all("img"):
            self._add_image(doc, img)

    def _picture_width(self, path: Path) -> Cm:
        """按比例限制：宽度不超过版心，高度不超过 9.5cm，避免一张长图吃掉整页。"""
        max_w, max_h = 13.5, 9.5
        try:
            from PIL import Image

            with Image.open(path) as im:
                w, h = im.size
            if w and h:
                width = min(max_w, max_h * w / h)
                return Cm(max(3.0, width))
        except Exception:  # noqa: BLE001
            pass
        return Cm(max_w)

    def _add_image(self, doc: Document, img: Tag) -> None:
        src = img.get("src") or ""
        if not src:
            return
        path = Path(src) if Path(src).is_absolute() else self.image_dir / Path(src).name
        if not path.exists() or path.stat().st_size == 0:
            return
        if path.suffix.lower() in {".svg", ".webp", ".avif"}:
            return  # Word 不支持，已在下载阶段尝试转 PNG
        try:
            doc.add_picture(str(path), width=self._picture_width(path))
            pic_para = doc.paragraphs[-1]
            pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_para.paragraph_format.space_after = Pt(4)
        except Exception:  # noqa: BLE001 - 单张图失败不影响整篇
            return
        alt = (img.get("alt") or "").strip()
        if alt and len(alt) < 120:
            para = doc.add_paragraph(alt)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(8)
            for run in para.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def _render_table(self, doc: Document, tag: Tag) -> None:
        rows = tag.find_all("tr")
        if not rows:
            return
        cols = max(len(r.find_all(["td", "th"])) for r in rows)
        if cols == 0:
            return
        table = doc.add_table(rows=0, cols=cols)
        table.style = "Light Grid Accent 1"
        table.autofit = True
        for tr in rows:
            cells = tr.find_all(["td", "th"])
            row = table.add_row().cells
            for i, cell in enumerate(cells[:cols]):
                row[i].text = cell.get_text(" ", strip=True)
                for para in row[i].paragraphs:
                    para.paragraph_format.space_after = Pt(1)
                    for run in para.runs:
                        run.font.size = Pt(8.5)
        doc.add_paragraph()


# --------------------------------------------------------------------------- markdown


def write_markdown(items: list[Item], path: Path, image_rel: str = "images", title: str = "链接归档") -> Path:
    ok_count = sum(1 for it in items if it.status == "ok")
    lines = [
        f"# {title}",
        "",
        f"> 生成时间：{_stamp()}　共 {len(items)} 条链接，成功 {ok_count} 条，异常 {len(items) - ok_count} 条。",
        "",
        "## 目录",
        "",
    ]
    for idx, item in enumerate(items, 1):
        name = (item.title or item.url).replace("[", "［").replace("]", "］")
        lines.append(f"{idx}. [{name}](#{idx})　`{item.date or '时间未知'}`")
    lines.append("")

    for idx, item in enumerate(items, 1):
        lines.append(f'<a id="{idx}"></a>')
        lines.append("")
        lines.append(f"## {idx}. {item.title or item.url}")
        lines.append("")
        lines.append(f"- **链接**：{item.url}")
        lines.append(f"- **时间**：{item.date or '未识别'}")
        if item.author:
            lines.append(f"- **作者**：{item.author}")
        if item.sitename:
            lines.append(f"- **站点**：{item.sitename}")
        if item.summary:
            lines.append(f"- **摘要**：{item.summary}")
        if item.error:
            lines.append(f"- **提示**：⚠ {item.error}")
        lines.append("")
        if item.content_html:
            lines.append(_html_to_markdown(item.content_html, image_rel))
        elif item.status == "failed":
            lines.append("_（抓取失败，未获取到正文）_")
        lines.append("")
        lines.append("---")
        lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _html_to_markdown(body_html: str, image_rel: str) -> str:
    from markdownify import MarkdownConverter

    class _Converter(MarkdownConverter):
        def convert_img(self, el, text, parent_tags=None):
            src = el.get("src") or ""
            if not src:
                return ""
            if not src.startswith(("http://", "https://", "data:")):
                src = f"{image_rel}/{Path(src).name}"
            alt = (el.get("alt") or "").strip() or "图片"
            return f"\n\n![{alt}]({src})\n\n"

    return _Converter(heading_style="atx", bullets="-").convert(body_html).strip()


# --------------------------------------------------------------------------- xlsx


def write_xlsx(items: list[Item], path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "抓取结果"
    headers = ["序号", "标题", "作者", "站点", "发布时间", "链接", "摘要", "正文", "正文长度", "抓取方式", "状态", "错误"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F3B57")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, item in enumerate(items, 1):
        ws.append(
            [
                idx,
                item.title,
                item.author,
                item.sitename,
                item.date,
                item.url,
                item.summary,
                item.content_text,
                len(item.content_text),
                item.fetch_mode,
                "OK" if item.status == "ok" else ("部分" if item.status == "partial" else "失败"),
                item.error,
            ]
        )

    widths = {1: 6, 2: 42, 3: 14, 4: 18, 5: 13, 6: 46, 7: 46, 8: 70, 9: 9, 10: 10, 11: 8, 12: 30}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- json / 失败清单


def write_json(items: list[Item], path: Path) -> Path:
    payload = [
        {
            "title": it.title,
            "author": it.author,
            "sitename": it.sitename,
            "date": it.date,
            "url": it.url,
            "summary": it.summary,
            "content_text": it.content_text,
            "content_html": it.content_html,
            "images": it.images,
            "fetch_mode": it.fetch_mode,
            "status": it.status,
            "error": it.error,
        }
        for it in items
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def write_failed(items: list[Item], out_dir: Path) -> list[Path]:
    bad = failed_items(items)
    if not bad:
        return []
    out_dir.mkdir(parents=True, exist_ok=True)

    md = out_dir / "failed.md"
    lines = [f"# 抓取失败清单（{_stamp()}）", "", f"共 {len(bad)} 条：", ""]
    for it in bad:
        lines.append(f"- {it.url}")
        lines.append(f"  - 状态：{it.status}　原因：{it.error or '未知'}")
    md.write_text("\n".join(lines), encoding="utf-8")

    csv_path = out_dir / "failed.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["链接", "标题", "状态", "错误原因", "抓取方式"])
        for it in bad:
            writer.writerow([it.url, it.title, it.status, it.error, it.fetch_mode])
    return [md, csv_path]


def export_all(items: list[Item], out_dir: Path, image_dir: Path, opts: Options, base_name: str) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    produced: dict[str, Path] = {}
    title = f"链接归档 {datetime.now().strftime('%Y-%m-%d')}"
    base = _safe_filename(base_name)
    if "docx" in opts.formats:
        produced["docx"] = DocxWriter(image_dir, opts).build(items, out_dir / f"{base}.docx", title)
    if "md" in opts.formats:
        produced["md"] = write_markdown(items, out_dir / f"{base}.md", image_rel=image_dir.name, title=title)
    if "xlsx" in opts.formats:
        produced["xlsx"] = write_xlsx(items, out_dir / f"{base}.xlsx")
    if "json" in opts.formats:
        produced["json"] = write_json(items, out_dir / f"{base}.json")
    for path in write_failed(items, out_dir):
        produced[f"failed_{path.suffix.lstrip('.')}"] = path
    return produced
